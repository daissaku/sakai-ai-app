"""
社内ナレッジ検索型 AIコンサルティングWebアプリ（RAGシステム）
Streamlit + ChromaDB + Gemini API
"""

import io
import os
import sqlite3
import uuid
from datetime import datetime
from pathlib import Path

import streamlit as st
import streamlit.components.v1 as components
from dotenv import load_dotenv
from google import genai
from google.genai import types
import chromadb
from pypdf import PdfReader
from pptx import Presentation

load_dotenv()

APP_PASSWORD = os.getenv("APP_PASSWORD", "")
GEMINI_API_KEY = os.getenv("GEMINI_API_KEY", "")
GEMINI_MODEL = os.getenv("GEMINI_MODEL", "gemini-flash-latest")
EMBEDDING_MODEL = "gemini-embedding-001"

BASE_DIR = Path(__file__).parent
DATA_DIR = BASE_DIR / "data"
DB_DIR = BASE_DIR / "chroma_db"
CHAT_DB_PATH = BASE_DIR / "chat_history.db"
COLLECTION_NAME = "sakai_knowledge"

MAX_VISIBLE_HISTORY = 10
IMAGE_MIME = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg"}

SYSTEM_PROMPT = (
    "あなたはコンサルタントの酒井です。"
    "提供された過去の企画書やメモの文脈に基づき、"
    "酒井の思考・価値観・ロジックを反映させてアドバイスしてください。"
)

st.set_page_config(page_title="SAKAIのAI", page_icon="💡", layout="wide")


# ---------------------------------------------------------------------------
# 認証
# ---------------------------------------------------------------------------
def check_password() -> bool:
    if st.session_state.get("authenticated"):
        return True

    st.title("🔒 SAKAIのAI ログイン")

    if not APP_PASSWORD:
        st.error(".env に APP_PASSWORD が設定されていません。管理者に確認してください。")
        return False

    with st.form("login_form"):
        pw = st.text_input("パスワードを入力してください", type="password")
        submitted = st.form_submit_button("ログイン")

    if submitted:
        if pw == APP_PASSWORD:
            st.session_state.authenticated = True
            st.rerun()
        else:
            st.error("パスワードが違います。")

    return False


# ---------------------------------------------------------------------------
# チャット履歴DB（SQLite）
# ---------------------------------------------------------------------------
def get_db_connection() -> sqlite3.Connection:
    conn = sqlite3.connect(str(CHAT_DB_PATH))
    conn.row_factory = sqlite3.Row
    return conn


def init_chat_db() -> None:
    conn = get_db_connection()
    conn.execute(
        """CREATE TABLE IF NOT EXISTS chats (
            id TEXT PRIMARY KEY,
            title TEXT NOT NULL,
            created_at TEXT NOT NULL,
            updated_at TEXT NOT NULL
        )"""
    )
    conn.execute(
        """CREATE TABLE IF NOT EXISTS messages (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            chat_id TEXT NOT NULL,
            role TEXT NOT NULL,
            content TEXT NOT NULL,
            created_at TEXT NOT NULL
        )"""
    )
    conn.commit()
    conn.close()


def make_title(text: str, limit: int = 24) -> str:
    text = text.strip().replace("\n", " ")
    return text[:limit] + ("…" if len(text) > limit else "")


def create_chat(title: str) -> str:
    chat_id = str(uuid.uuid4())
    now = datetime.now().isoformat()
    conn = get_db_connection()
    conn.execute(
        "INSERT INTO chats (id, title, created_at, updated_at) VALUES (?, ?, ?, ?)",
        (chat_id, title, now, now),
    )
    conn.commit()
    conn.close()
    return chat_id


def touch_chat(chat_id: str) -> None:
    conn = get_db_connection()
    conn.execute(
        "UPDATE chats SET updated_at = ? WHERE id = ?", (datetime.now().isoformat(), chat_id)
    )
    conn.commit()
    conn.close()


def add_message(chat_id: str, role: str, content: str) -> None:
    conn = get_db_connection()
    conn.execute(
        "INSERT INTO messages (chat_id, role, content, created_at) VALUES (?, ?, ?, ?)",
        (chat_id, role, content, datetime.now().isoformat()),
    )
    conn.commit()
    conn.close()
    touch_chat(chat_id)


def get_messages(chat_id: str) -> list[dict]:
    conn = get_db_connection()
    rows = conn.execute(
        "SELECT role, content FROM messages WHERE chat_id = ? ORDER BY id ASC", (chat_id,)
    ).fetchall()
    conn.close()
    return [{"role": r["role"], "content": r["content"]} for r in rows]


def list_chats(search: str = "") -> list[dict]:
    conn = get_db_connection()
    if search.strip():
        like = f"%{search.strip()}%"
        rows = conn.execute(
            """SELECT DISTINCT c.id, c.title, c.updated_at FROM chats c
               LEFT JOIN messages m ON m.chat_id = c.id
               WHERE c.title LIKE ? OR m.content LIKE ?
               ORDER BY c.updated_at DESC""",
            (like, like),
        ).fetchall()
    else:
        rows = conn.execute("SELECT id, title, updated_at FROM chats ORDER BY updated_at DESC").fetchall()
    conn.close()
    return [{"id": r["id"], "title": r["title"], "updated_at": r["updated_at"]} for r in rows]


def delete_chat(chat_id: str) -> None:
    conn = get_db_connection()
    conn.execute("DELETE FROM messages WHERE chat_id = ?", (chat_id,))
    conn.execute("DELETE FROM chats WHERE id = ?", (chat_id,))
    conn.commit()
    conn.close()


# ---------------------------------------------------------------------------
# ドキュメント読み込み・チャンク分割（社内ナレッジDB用）
# ---------------------------------------------------------------------------
def read_txt(path: Path) -> str:
    for enc in ("utf-8", "cp932", "shift_jis"):
        try:
            return path.read_text(encoding=enc)
        except (UnicodeDecodeError, LookupError):
            continue
    return path.read_text(encoding="utf-8", errors="ignore")


def read_pdf(path: Path) -> str:
    reader = PdfReader(str(path))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def load_documents() -> list[dict]:
    docs = []
    if not DATA_DIR.exists():
        return docs
    for path in sorted(DATA_DIR.iterdir()):
        if not path.is_file():
            continue
        suffix = path.suffix.lower()
        if suffix == ".pdf":
            text = read_pdf(path)
        elif suffix in (".txt", ".md"):
            text = read_txt(path)
        else:
            continue
        if text.strip():
            docs.append({"source": path.name, "text": text})
    return docs


def chunk_text(text: str, chunk_size: int = 800, overlap: int = 100) -> list[str]:
    chunks = []
    start = 0
    length = len(text)
    while start < length:
        end = min(start + chunk_size, length)
        chunk = text[start:end].strip()
        if chunk:
            chunks.append(chunk)
        if end == length:
            break
        start = end - overlap
    return chunks


# ---------------------------------------------------------------------------
# アップロードファイル（画像・PDF・PowerPoint）の解析
# ---------------------------------------------------------------------------
def extract_pdf_text_from_bytes(data: bytes) -> str:
    reader = PdfReader(io.BytesIO(data))
    return "\n".join(page.extract_text() or "" for page in reader.pages)


def extract_pptx_text_from_bytes(data: bytes) -> str:
    prs = Presentation(io.BytesIO(data))
    slide_texts = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape.text_frame.text.strip():
                parts.append(shape.text_frame.text)
        if parts:
            slide_texts.append(f"[スライド{i}]\n" + "\n".join(parts))
    return "\n\n".join(slide_texts)


def build_uploaded_content(uploaded_files) -> tuple[list, str]:
    """アップロードされたファイルから (画像Partのリスト, 抽出テキスト) を返す"""
    image_parts = []
    texts = []
    for f in uploaded_files or []:
        ext = Path(f.name).suffix.lower()
        data = f.getvalue()
        try:
            if ext in IMAGE_MIME:
                image_parts.append(types.Part.from_bytes(data=data, mime_type=IMAGE_MIME[ext]))
            elif ext == ".pdf":
                text = extract_pdf_text_from_bytes(data)
                if text.strip():
                    texts.append(f"[アップロードファイル: {f.name}]\n{text}")
            elif ext == ".pptx":
                text = extract_pptx_text_from_bytes(data)
                if text.strip():
                    texts.append(f"[アップロードファイル: {f.name}]\n{text}")
        except Exception as e:
            texts.append(f"[アップロードファイル: {f.name}] 読み込みエラー: {e}")
    return image_parts, "\n\n".join(texts)


# ---------------------------------------------------------------------------
# Gemini クライアント・埋め込み
# ---------------------------------------------------------------------------
@st.cache_resource(show_spinner=False)
def get_genai_client() -> genai.Client:
    return genai.Client(api_key=GEMINI_API_KEY)


def embed_texts(texts: list[str], task_type: str) -> list[list[float]]:
    client = get_genai_client()
    result = client.models.embed_content(
        model=EMBEDDING_MODEL,
        contents=texts,
        config=types.EmbedContentConfig(task_type=task_type),
    )
    return [e.values for e in result.embeddings]


# ---------------------------------------------------------------------------
# ベクトルDB（ChromaDB）
# ---------------------------------------------------------------------------
@st.cache_resource(show_spinner=False)
def get_chroma_client():
    DB_DIR.mkdir(parents=True, exist_ok=True)
    return chromadb.PersistentClient(path=str(DB_DIR))


def get_or_create_collection():
    client = get_chroma_client()
    return client.get_or_create_collection(
        name=COLLECTION_NAME, metadata={"hnsw:space": "cosine"}
    )


def rebuild_database() -> int:
    client = get_chroma_client()
    try:
        client.delete_collection(COLLECTION_NAME)
    except Exception:
        pass
    collection = client.get_or_create_collection(
        name=COLLECTION_NAME, metadata={"hnsw:space": "cosine"}
    )

    docs = load_documents()
    if not docs:
        return 0

    all_chunks, all_ids, all_metas = [], [], []
    for doc in docs:
        for i, chunk in enumerate(chunk_text(doc["text"])):
            all_chunks.append(chunk)
            all_ids.append(f"{doc['source']}::{i}")
            all_metas.append({"source": doc["source"]})

    batch_size = 50
    for i in range(0, len(all_chunks), batch_size):
        batch_chunks = all_chunks[i : i + batch_size]
        batch_embeddings = embed_texts(batch_chunks, task_type="RETRIEVAL_DOCUMENT")
        collection.add(
            ids=all_ids[i : i + batch_size],
            documents=batch_chunks,
            metadatas=all_metas[i : i + batch_size],
            embeddings=batch_embeddings,
        )

    return len(all_chunks)


def search_context(query: str, top_k: int = 5):
    collection = get_or_create_collection()
    count = collection.count()
    if count == 0:
        return [], []
    query_embedding = embed_texts([query], task_type="RETRIEVAL_QUERY")[0]
    results = collection.query(
        query_embeddings=[query_embedding], n_results=min(top_k, count)
    )
    documents = results["documents"][0] if results["documents"] else []
    metadatas = results["metadatas"][0] if results["metadatas"] else []
    return documents, metadatas


# ---------------------------------------------------------------------------
# Gemini 回答生成
# ---------------------------------------------------------------------------
def generate_answer(question: str, history: list[dict], uploaded_files=None) -> tuple[str, list[str]]:
    documents, metadatas = search_context(question)
    sources = sorted({m["source"] for m in metadatas}) if metadatas else []

    if documents:
        context_text = "\n\n---\n\n".join(
            f"[{m['source']}]\n{d}" for d, m in zip(documents, metadatas)
        )
    else:
        context_text = "（関連する過去資料は見つかりませんでした）"

    convo = ""
    for turn in history[-6:]:
        role = "ユーザー" if turn["role"] == "user" else "酒井"
        convo += f"{role}: {turn['content']}\n"

    image_parts, uploaded_text = build_uploaded_content(uploaded_files)

    prompt_text = (
        f"### 過去資料からの関連文脈\n{context_text}\n\n"
        f"### これまでの会話\n{convo if convo else '（なし）'}\n"
    )
    if uploaded_text:
        prompt_text += f"### アップロードされたファイルの内容\n{uploaded_text}\n\n"
    prompt_text += f"### 今回の質問\n{question}\n\n上記の文脈を踏まえ、酒井として回答してください。"

    contents = [prompt_text, *image_parts]

    client = get_genai_client()
    response = client.models.generate_content(
        model=GEMINI_MODEL,
        contents=contents,
        config=types.GenerateContentConfig(system_instruction=SYSTEM_PROMPT),
    )
    return response.text, sources


# ---------------------------------------------------------------------------
# 回答表示後のスクロール制御（回答の先頭にスクロール）
# ---------------------------------------------------------------------------
def scroll_to_top_of_answer(anchor_id: str) -> None:
    components.html(
        f"""
        <script>
        (function() {{
            const scrollToAnchor = () => {{
                const doc = window.parent.document;
                const el = doc.getElementById("{anchor_id}");
                if (el) {{
                    el.scrollIntoView({{behavior: "smooth", block: "start"}});
                    return true;
                }}
                return false;
            }};
            if (!scrollToAnchor()) {{
                let attempts = 0;
                const timer = setInterval(() => {{
                    attempts += 1;
                    if (scrollToAnchor() || attempts > 20) {{
                        clearInterval(timer);
                    }}
                }}, 100);
            }}
        }})();
        </script>
        """,
        height=0,
    )


# ---------------------------------------------------------------------------
# サイドバー（Gemini / ChatGPT 風）
# ---------------------------------------------------------------------------
def render_chat_row(chat: dict) -> None:
    col1, col2 = st.columns([5, 1])
    with col1:
        label = chat["title"] or "(無題のチャット)"
        is_active = chat["id"] == st.session_state.current_chat_id
        if st.button(
            label,
            key=f"load_{chat['id']}",
            use_container_width=True,
            type="primary" if is_active else "secondary",
        ):
            st.session_state.current_chat_id = chat["id"]
            st.session_state.uploader_key += 1
            st.rerun()
    with col2:
        if st.button("🗑️", key=f"del_{chat['id']}", help="この履歴を削除"):
            delete_chat(chat["id"])
            if st.session_state.current_chat_id == chat["id"]:
                st.session_state.current_chat_id = None
            st.rerun()


@st.dialog("過去のチャット履歴")
def show_all_history_dialog(chats: list[dict]) -> None:
    with st.container(height=420):
        for chat in chats:
            render_chat_row(chat)


def render_sidebar() -> None:
    with st.sidebar:
        st.markdown("### 💡 SAKAIのAI")

        # ① チャットを新規作成
        if st.button("✏️ チャットを新規作成", use_container_width=True, type="primary"):
            st.session_state.current_chat_id = None
            st.session_state.uploader_key += 1
            st.rerun()

        # ② チャットを検索
        search_query = st.text_input(
            "チャットを検索",
            key="history_search",
            placeholder="🔍 チャットを検索",
            label_visibility="collapsed",
        )

        st.divider()

        # ③ 過去のチャット履歴リスト
        chats = list_chats(search=search_query)
        st.caption(f"履歴（{len(chats)}件）" if search_query else "履歴")

        if not chats:
            st.caption("まだチャット履歴がありません。")
        else:
            for chat in chats[:MAX_VISIBLE_HISTORY]:
                render_chat_row(chat)

            if len(chats) > MAX_VISIBLE_HISTORY:
                if st.button("過去のチャットを確認", use_container_width=True):
                    show_all_history_dialog(chats)

        st.divider()

        # ④ 設定セクション
        st.markdown("##### ⚙️ 設定")

        file_count = (
            len([p for p in DATA_DIR.iterdir() if p.is_file()]) if DATA_DIR.exists() else 0
        )
        st.caption(f"data/ 内のファイル数: {file_count}")

        collection = get_or_create_collection()
        st.caption(f"インデックス済みチャンク数: {collection.count()}")

        if "rebuild_message" in st.session_state:
            msg_kind, msg_text = st.session_state.pop("rebuild_message")
            getattr(st, msg_kind)(msg_text)

        if st.button("🔄 データベース再構築", use_container_width=True):
            if not DATA_DIR.exists() or file_count == 0:
                st.session_state.rebuild_message = ("warning", "data/ フォルダにファイルがありません。")
            else:
                with st.spinner("data/ 内のドキュメントを読み込み、ベクトルDBを再構築しています..."):
                    try:
                        n = rebuild_database()
                        st.session_state.rebuild_message = ("success", f"再構築が完了しました（{n} チャンク）")
                    except Exception as e:
                        st.session_state.rebuild_message = ("error", f"再構築中にエラーが発生しました: {e}")
            st.rerun()

        st.caption(f"📁 `{DATA_DIR}`")
        if st.button("📂 データフォルダを開く", use_container_width=True):
            try:
                os.startfile(str(DATA_DIR))  # Windows専用
            except Exception as e:
                st.error(f"フォルダを開けませんでした: {e}")

        # ⑤ ログアウト（一番下）
        st.divider()
        if st.button("🚪 ログアウト", use_container_width=True):
            st.session_state.authenticated = False
            st.rerun()


# ---------------------------------------------------------------------------
# メインUI
# ---------------------------------------------------------------------------
def main_app():
    init_chat_db()

    if "current_chat_id" not in st.session_state:
        st.session_state.current_chat_id = None
    if "uploader_key" not in st.session_state:
        st.session_state.uploader_key = 0

    st.title("💡 SAKAIのAI")
    st.caption(f"モデル: {GEMINI_MODEL} / 埋め込み: {EMBEDDING_MODEL}")

    messages = get_messages(st.session_state.current_chat_id) if st.session_state.current_chat_id else []
    for msg in messages:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])

    uploaded_files = st.file_uploader(
        "ファイルを添付（画像 / PDF / PowerPoint）",
        type=["png", "jpg", "jpeg", "pdf", "pptx"],
        accept_multiple_files=True,
        key=f"uploader_{st.session_state.uploader_key}",
    )

    question = st.chat_input("質問を入力してください（例：新規事業の企画で気をつけるべき点は？）")
    if question:
        if st.session_state.current_chat_id is None:
            st.session_state.current_chat_id = create_chat(make_title(question))
        chat_id = st.session_state.current_chat_id

        history_for_context = get_messages(chat_id)
        add_message(chat_id, "user", question)
        with st.chat_message("user"):
            st.markdown(question)

        anchor_id = f"answer-anchor-{uuid.uuid4().hex}"
        with st.chat_message("assistant"):
            st.markdown(f'<div id="{anchor_id}"></div>', unsafe_allow_html=True)
            with st.spinner("酒井として考えています..."):
                try:
                    answer, sources = generate_answer(question, history_for_context, uploaded_files)
                    if sources:
                        answer_full = answer + "\n\n---\n**参照した過去資料:** " + "、".join(sources)
                    else:
                        answer_full = answer + "\n\n---\n**参照した過去資料:** なし"
                except Exception as e:
                    answer_full = f"エラーが発生しました: {e}"
            st.markdown(answer_full)

        add_message(chat_id, "assistant", answer_full)
        st.session_state.uploader_key += 1
        scroll_to_top_of_answer(anchor_id)

    # サイドバーはチャット処理の後に描画し、最新の履歴を反映させる
    render_sidebar()


def main():
    if not GEMINI_API_KEY:
        st.error(".env に GEMINI_API_KEY が設定されていません。.env.example を参考に設定してください。")
        st.stop()

    if check_password():
        main_app()


if __name__ == "__main__":
    main()
